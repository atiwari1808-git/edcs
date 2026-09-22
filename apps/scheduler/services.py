"""Booking queries, stats, filters, and Excel export for the dashboard."""
import io
from datetime import date
from django.db.models import Q
from django.contrib.auth import get_user_model
from .models import Meeting

ACTIVE = ["PENDING", "REQUESTED", "CONFIRMED"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _display_status(m: Meeting, today=None) -> str:
    today = today or date.today()
    if m.status == "CANCELLED":
        return "cancelled"
    if m.status == "FAILED":
        return "failed"
    d = m.booking_date or m.start_at.date()
    return "completed" if d < today else "scheduled"


def _resolve_scrum_master_names(emails):
    """Map a set of scrum-master emails to a display name from the Users DB.
    Falls back to the raw email when no matching user exists."""
    emails = {e.lower() for e in emails if e}
    if not emails:
        return {}
    User = get_user_model()
    name_map = {}
    for u in User.objects.filter(email__in=emails):
        full = (u.get_full_name() or "").strip()
        name_map[u.email.lower()] = full or u.username
    # anything not found falls back to the email itself
    return {e: name_map.get(e, e) for e in emails}


# ---------------------------------------------------------------------------
# Querysets
# ---------------------------------------------------------------------------
def all_bookings():
    return (
        Meeting.objects
        .select_related("organizer", "slot", "validation_run",
                        "validation_run__automation_type",
                        "validation_run__jira_type")
        .prefetch_related("attendees")
        .order_by("-start_at")
    )


def filtered_bookings(q="", tool="", status="", date_from="", date_to="",
                      jira_type="", sm_scope_email=""):
    rows = all_bookings()
    if q:
        rows = rows.filter(
            Q(organizer__username__icontains=q) | Q(tool_key__icontains=q) |
            Q(subject__icontains=q) | Q(validation_run__jira_id__icontains=q) |
            Q(validation_run__parent_jira_id__icontains=q) |
            Q(attendees__email__icontains=q)
        ).distinct()
    if tool:
        rows = rows.filter(tool_key=tool)
    if jira_type:
        rows = rows.filter(validation_run__jira_type__key=jira_type)
    if sm_scope_email:
        # Scrum-master scope: only meetings where this email is a Scrum Master
        rows = rows.filter(attendees__email__iexact=sm_scope_email,
                           attendees__type="SCRUM_MASTER").distinct()
    if date_from:
        rows = rows.filter(booking_date__gte=date_from)
    if date_to:
        rows = rows.filter(booking_date__lte=date_to)

    rows = list(rows)

    # Bulk-resolve all scrum-master emails to names in a single query.
    all_sm_emails = set()
    for m in rows:
        for a in m.attendees.all():
            if a.type == "SCRUM_MASTER":
                all_sm_emails.add(a.email.lower())
    sm_names = _resolve_scrum_master_names(all_sm_emails)

    out = []
    for m in rows:
        disp = _display_status(m)
        if status and disp != status:
            continue
        vr = m.validation_run
        at = vr.automation_type if vr else None
        jt = vr.jira_type if vr else None
        scrum_masters = [
            {"email": a.email, "name": sm_names.get(a.email.lower(), a.email)}
            for a in m.attendees.all() if a.type == "SCRUM_MASTER"
        ]
        out.append({
            "m": m,
            "display_status": disp,
            "jira_id": vr.jira_id if vr else "",
            "customer_name": vr.customer_name if vr else "",
            "automation_name": vr.automation_name if vr else "",
            "automation_type_display": at.display_name if at else "—",
            "automation_type_key": at.key if at else "",
            "requires_scheduling": at.requires_scheduling if at else True,
            # JIRA Type / enhancement (Req 2)
            "jira_type_display": jt.display_name if jt else "—",
            "jira_type_key": jt.key if jt else "",
            "parent_jira_id": vr.parent_jira_id if vr else "",
            "enhancement_jira_id": vr.enhancement_jira_id if vr else "",
            # Scrum masters (Req 5)
            "scrum_masters": scrum_masters,
        })
    return out


# ---------------------------------------------------------------------------
# Stats
# ---------------------------------------------------------------------------
def booking_stats():
    today = date.today()
    rows = list(Meeting.objects.all())
    return {
        "total":     len(rows),
        "upcoming":  sum(1 for m in rows if m.status in ACTIVE
                         and (m.booking_date or m.start_at.date()) >= today),
        "completed": sum(1 for m in rows if m.status in ACTIVE
                         and (m.booking_date or m.start_at.date()) < today),
        "cancelled": sum(1 for m in rows if m.status == "CANCELLED"),
    }


def no_meeting_stats():
    """Count of doc-only (Healthcheck / Backup) PASSED runs — for dashboard stat card."""
    from apps.validation.models import ValidationRun
    return ValidationRun.objects.filter(
        automation_type__requires_scheduling=False,
        status="PASSED",
    ).count()


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
_CLR = {
    "header_fill":   "FF1C3C5A",
    "header_font":   "FFFFFFFF",
    "subheader":     "FFD9E8F5",
    "row_even":      "FFF5F9FF",
    "row_odd":       "FFFFFFFF",
    "type_exec":     "FFD6EAD7",
    "type_hc":       "FFFFF3CD",
    "type_backup":   "FFE2D9F3",
    "type_other":    "FFEDEDED",
    "pass_green":    "FFD6EAD7",
    "fail_red":      "FFFFD7D7",
}
_TYPE_COLOURS = {
    "EXECUTION":   _CLR["type_exec"],
    "HEALTHCHECK": _CLR["type_hc"],
    "BACKUP":      _CLR["type_backup"],
}


def _type_colour(key: str) -> str:
    return _TYPE_COLOURS.get((key or "").upper(), _CLR["type_other"])


def _status_colour(status: str) -> str:
    s = (status or "").lower()
    if s in ("scheduled", "completed"):
        return _CLR["pass_green"]
    if s in ("cancelled", "failed"):
        return _CLR["fail_red"]
    return _CLR["type_other"]

def bookings_workbook(rows) -> io.BytesIO:
    """Build a styled two-sheet Excel workbook with ALL handover details."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from apps.validation.models import ValidationRun

    wb = Workbook()

    def _fill(hex_argb):   return PatternFill("solid", fgColor=hex_argb)
    def _border():
        s = Side(style="thin", color="FFD0D0D0")
        return Border(left=s, right=s, top=s, bottom=s)
    def _header_font():    return Font(bold=True, color=_CLR["header_font"], size=10)
    def _data_font(bold=False): return Font(bold=bold, size=10)
    def _center():         return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def _left():           return Alignment(horizontal="left",  vertical="center", wrap_text=True)

    def _style_header_row(ws, row_idx=1):
        for cell in ws[row_idx]:
            cell.fill = _fill(_CLR["header_fill"]); cell.font = _header_font()
            cell.alignment = _center(); cell.border = _border()

    def _auto_width(ws, min_w=12, max_w=48):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max(min_w, min(max_w, max(len(str(c.value or "")) for c in col) + 3))
            ws.column_dimensions[letter].width = width

    def _dt(v):  return v.strftime("%Y-%m-%d %H:%M") if v else ""

    # ================= SHEET 1 — Handover Bookings ================= #
    ws1 = wb.active
    ws1.title = "Handover Bookings"
    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 22
    HEADERS_S1 = [
        "Organizer", "Tool", "Automation Type", "JIRA Type", "Customer",
        "JIRA ID", "Parent JIRA", "Enhancement JIRA", "Automation Name",
        "Scrum Masters", "Developers", "CC",
        "Handover Initiated", "Handover Date", "Slot",
        "Status", "Cancelled / Rescheduled By", "Reason",
    ]
    ws1.append(HEADERS_S1)
    _style_header_row(ws1)
    AT_COL_1, STATUS_COL_1 = 3, 16
    CENTER_1 = {2, 3, 4, 6, 7, 8, 13, 14, 15, 16}

    for idx, r in enumerate(rows, start=2):
        m    = r["m"]
        att  = list(m.attendees.all())
        pick = lambda k: ", ".join(a.email for a in att if a.type == k)
        sm_names = ", ".join(s["name"] for s in r.get("scrum_masters", [])) or pick("SCRUM_MASTER")
        at_key  = r.get("automation_type_key", "")
        at_disp = r.get("automation_type_display", "")
        disp_st = r["display_status"]
        cancelled_by = m.cancelled_by.username if m.cancelled_by else ""
        ws1.append([
            m.organizer.username,
            m.tool_key,
            at_disp,
            r.get("jira_type_display", ""),
            m.validation_run.customer_name if m.validation_run else "",
            r["jira_id"],
            r.get("parent_jira_id", ""),
            r.get("enhancement_jira_id", ""),
            r["automation_name"],
            sm_names,
            pick("DEVELOPER"),
            pick("CC"),
            _dt(m.created_at),                       # Handover Initiated
            _dt(m.start_at),                         # Handover Date (date+time)
            m.slot.label if m.slot else "",
            disp_st,
            cancelled_by,
            m.cancellation_reason or "",
        ])
        row_fill = _fill(_CLR["row_even"] if idx % 2 == 0 else _CLR["row_odd"])
        for col_idx, cell in enumerate(ws1[idx], start=1):
            cell.border = _border()
            cell.alignment = _center() if col_idx in CENTER_1 else _left()
            cell.font = _data_font()
            if col_idx == AT_COL_1:      cell.fill = _fill(_type_colour(at_key))
            elif col_idx == STATUS_COL_1: cell.fill = _fill(_status_colour(disp_st))
            else:                         cell.fill = row_fill
    _auto_width(ws1)

    # ================= SHEET 2 — No Meeting Needed ================= #
    ws2 = wb.create_sheet("No Meeting Needed")
    ws2.freeze_panes = "A2"
    ws2.row_dimensions[1].height = 22
    HEADERS_S2 = [
        "Organizer", "Tool", "Automation Type", "JIRA Type", "Customer",
        "JIRA ID", "Parent JIRA", "Enhancement JIRA", "Automation Name",
        "Status", "Remarks", "Handover Initiated", "Handover Date",
    ]
    ws2.append(HEADERS_S2)
    _style_header_row(ws2)
    AT_COL_2, STATUS_COL_2 = 3, 10
    CENTER_2 = {2, 3, 4, 6, 7, 8, 10, 12, 13}

    no_meeting_runs = (
        ValidationRun.objects
        .filter(automation_type__requires_scheduling=False)
        .select_related("requested_by", "tool", "automation_type", "jira_type")
        .order_by("-started_at")
    )
    STATUS_REMARKS = {
        "PASSED": "No meeting needed - verification passed",
        "FAILED": "Verification failed - no meeting required",
        "CANCELLED": "Scan cancelled", "RUNNING": "Scan in progress",
        "PENDING": "Scan pending", "ERROR": "Scan error",
    }
    for idx, vr in enumerate(no_meeting_runs, start=2):
        at = vr.automation_type
        at_key = at.key if at else ""; at_disp = at.display_name if at else ""
        jt = vr.jira_type
        ws2.append([
            vr.requested_by.username if vr.requested_by else "",
            vr.tool.key if vr.tool else "",
            at_disp,
            jt.display_name if jt else "",
            getattr(vr, "customer_name", ""),
            vr.jira_id,
            getattr(vr, "parent_jira_id", ""),
            getattr(vr, "enhancement_jira_id", ""),
            getattr(vr, "automation_name", ""),
            vr.status,
            STATUS_REMARKS.get(vr.status, vr.status),
            _dt(vr.started_at),                      # Handover Initiated
            _dt(vr.started_at),                      # Handover Date (mirrors initiated - no meeting)
        ])
        row_fill = _fill(_CLR["row_even"] if idx % 2 == 0 else _CLR["row_odd"])
        for col_idx, cell in enumerate(ws2[idx], start=1):
            cell.border = _border(); cell.font = _data_font()
            cell.alignment = _center() if col_idx in CENTER_2 else _left()
            if col_idx == AT_COL_2:       cell.fill = _fill(_type_colour(at_key))
            elif col_idx == STATUS_COL_2: cell.fill = _fill(_CLR["pass_green"] if vr.status == "PASSED" else _CLR["fail_red"])
            else:                         cell.fill = row_fill
    _auto_width(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf

